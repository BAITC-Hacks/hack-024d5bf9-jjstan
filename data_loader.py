"""Local, read-only adapters for the HackAlem 1C workbooks.

Public interface: load_data(paths: list[str]) -> dict[str, pandas.DataFrame].
No network, extracted files, UI, forecasting or implicit zero filling.
"""
from __future__ import annotations

from collections import defaultdict
from datetime import date, datetime
from hashlib import sha256
from io import BytesIO
from itertools import chain, islice
from pathlib import Path
import re
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZipFile

import numpy as np
import pandas as pd
from openpyxl import load_workbook

ALL = "*all*"  # Explicit user contract; TEAM_PLAN.md currently says __all__.
SCHEMAS = {
    "products": "supplier sku supplier_article name category unit min_order_qty pack_multiple".split(),
    "monthly_sales": "supplier sku month quantity is_complete".split(),
    "transactions": "supplier sku date document_id customer_id warehouse quantity transaction_type".split(),
    "stock": "supplier sku warehouse as_of free_stock is_current".split(),
    "transit": "supplier sku warehouse expected_date quantity".split(),
    "seasonality": "supplier category month_number factor".split(),
    "stockouts": "supplier sku warehouse start_date end_date".split(),
    "quality_report": "supplier sku severity issue source".split(),
}
DATE_COLUMNS = {"month", "date", "as_of", "expected_date", "start_date", "end_date", "report_month", "snapshot_as_of",
                "period_start", "period_end"}
NUMBER_COLUMNS = {"quantity", "free_stock", "min_order_qty", "pack_multiple", "factor",
                  "raw_quantity", "stock_quantity", "growth_coefficient_raw", "seasonality_coefficient_raw",
                  "reconciled_transaction_quantity", "reconciliation_difference"}
MONTHS = {v: i for i, v in enumerate(
    ["янв", "фев", "мар", "апр", "май", "июн", "июл", "авг", "сен", "окт", "ноя", "дек"], 1)}
SKU_HEADERS = {"код", "код 1с", "номенклатура.код"}


def _text(value):
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip() or None


def _label(value):
    return re.sub(r"\s+", " ", _text(value) or "").lower().replace("ё", "е")


def _month(value):
    if isinstance(value, (date, datetime)):
        return pd.Timestamp(value).replace(day=1).normalize()
    s = _label(value)
    m = re.fullmatch(r"([а-я]+)\.?\s+(20\d{2})(?:\s*г\.?)?", s)
    if m and m[1][:3] in MONTHS:
        return pd.Timestamp(int(m[2]), MONTHS[m[1][:3]], 1)
    return None


def _date(value):
    if value is None:
        return pd.NaT
    if isinstance(value, (datetime, date, pd.Timestamp)):
        return pd.Timestamp(value)
    if isinstance(value, str):
        value = value.strip()
        for fmt in ("%d.%m.%Y %H:%M:%S", "%d.%m.%Y", "%Y-%m-%d"):
            try:
                return pd.Timestamp(datetime.strptime(value, fmt))
            except ValueError:
                pass
        return pd.to_datetime(value, dayfirst=True, errors="coerce")
    return pd.NaT  # Do not guess the epoch of an unformatted numeric cell.


def _file_date(name):
    match = re.search(r"(?<!\d)(\d{2}\.\d{2}\.20\d{2})(?!\d)", name)
    return _date(match[1]) if match else pd.NaT


def _supplier(name):
    s = name.lower()
    if "system" in s or "syseme" in s:
        return "Systeme Electric"
    if re.search(r"iek|иэк", s):
        return "IEK"
    return None


def _kind(name):
    s = Path(name).name.lower()
    for key, value in [("динамика", "transactions"), ("moq", "moq"),
                       ("сезонность", "seasonality"), ("остатки", "history_stock"),
                       ("продажи", "monthly_sales"), ("пути", "snapshot"), ("путь", "snapshot")]:
        if key in s:
            return value
    return None


class _Loader:
    def __init__(self):
        self.rows = {name: [] for name in SCHEMAS}
        self.products = {}
        self.product_sources = defaultdict(set)
        self.units = defaultdict(set)
        self.suppliers = set()
        self.seen = set()
        self.issues = set()
        self.cutoffs = {}
        self.conflicts = defaultdict(set)

    def issue(self, supplier, sku, issue, source, severity="warning", **details):
        key = (supplier, sku, severity, issue, source)
        if key not in self.issues:
            self.issues.add(key)
            self.rows["quality_report"].append(dict(zip(SCHEMAS["quality_report"], key), **details))

    def number(self, value, supplier, sku, source):
        if value is None or _text(value) is None:
            return np.nan
        s = str(value).replace("\u00a0", "").replace(" ", "").replace(",", ".")
        try:
            n = float(s)
            if np.isfinite(n):
                return n
        except (TypeError, ValueError):
            pass
        self.issue(supplier, sku, "invalid_numeric_value", source, "error")
        return np.nan

    def product(self, supplier, sku, values, kind, source):
        key = (supplier, sku)
        p = self.products.setdefault(key, {"supplier": supplier, "sku": sku})
        self.product_sources[key].add(kind)
        for field, value in values.items():
            if value is None or (isinstance(value, float) and np.isnan(value)):
                continue
            if field == "unit":
                self.units[key].add(value)
            if field in p and p[field] != value and field != "name":
                self.issue(supplier, sku, f"conflicting_product_{field}", source)
                self.conflicts[key].add(field)
            else:
                p[field] = value

    def ingest(self, name, payload):
        supplier, kind = _supplier(name), _kind(name)
        digest = (supplier, sha256(payload).hexdigest())
        if digest in self.seen:
            self.issue(supplier, None, "duplicate_file_skipped", name, "info")
            return
        if supplier is None or kind is None:
            self.issue(supplier, None, "unrecognized_supplier_or_file_type", name, "error")
            return
        self.seen.add(digest)
        self.suppliers.add(supplier)
        cutoff = _file_date(Path(name).name)
        try:
            wb = load_workbook(BytesIO(payload), data_only=True, read_only=True)
        except (BadZipFile, OSError, ValueError, KeyError, ParseError) as exc:
            self.issue(supplier, None, f"unreadable_workbook:{type(exc).__name__}", name, "error")
            return
        parsed = False
        try:
            for ws in wb:
                source = f"{name}:{ws.title}"
                iterator = ws.iter_rows()
                head = list(islice(iterator, 20))
                if kind == "seasonality":
                    parsed |= self.seasonality(supplier, source, chain(head, iterator))
                    continue
                header = next((i for i, row in enumerate(head)
                               if any(_label(c.value) in SKU_HEADERS for c in row)), None)
                if header is None:
                    continue  # Embedded charts/seasonality are not another sales source.
                parsed = True
                headers = [_label(c.value) for c in head[header]]
                months = {i: _month(c.value) for i, c in enumerate(head[header]) if _month(c.value) is not None}
                meta = " ".join(_label(c.value) for row in head[header+1:header+3] for c in row)
                self.table(supplier, kind, source, headers, months,
                           chain(head[header+1:], iterator), cutoff, "нач. остаток" in meta)
        finally:
            wb.close()
        if not parsed:
            self.issue(supplier, None, "required_header_not_found", name, "error")
        elif kind == "snapshot" and pd.notna(cutoff):
            self.cutoffs[supplier] = max(cutoff, self.cutoffs.get(supplier, cutoff))

    def seasonality(self, supplier, source, rows):
        columns = None
        parsed = False
        seen_months = set()
        for row in rows:
            labels = [_label(c.value) for c in row]
            if "месяц" in labels and "сезонность" in labels:
                columns = labels.index("месяц"), labels.index("сезонность")
                continue
            if columns is None:
                continue
            a, b = columns
            month = MONTHS.get(labels[a][:3])
            if month:
                if month in seen_months:
                    break  # A second calculation block is not another supplier profile.
                seen_months.add(month)
                factor = self.number(row[b].value, supplier, None, source)
                if pd.isna(factor) or factor <= 0:
                    self.issue(supplier, None, "invalid_seasonality_factor", source, "error")
                    factor = np.nan
                self.rows["seasonality"].append(dict(supplier=supplier, category=ALL,
                    month_number=month, factor=factor, source=source))
                parsed = True
                if len(seen_months) == 12:
                    break
        if parsed:
            self.issue(supplier, None, "supplier_seasonality_includes_partial_2026_and_unconfirmed_measure", source)
        return parsed

    def table(self, supplier, kind, source, headers, months, rows, cutoff, opening):
        def idx(*names):
            return next((headers.index(x) for x in names if x in headers), None)
        code = next(i for i,h in enumerate(headers) if h in SKU_HEADERS)
        name_col = idx("номенклатура", "наименование")
        unit_col = idx("ед.", "ед.изм", "ед. изм.")
        article_col = idx("артикул поставщика", "артикул иэк", "артикул")
        month_latest = max(months.values()) if months else pd.NaT
        transit_cols = {i: h for i,h in enumerate(headers) if "поступление до" in h or "сэ в пути" in h}
        if kind in {"monthly_sales", "history_stock"} and not months:
            self.issue(supplier, None, "month_headers_missing", source, "error")
        if kind == "transactions":
            for field in ["дата", "номер", "документ", "количество"]:
                if idx(field) is None:
                    self.issue(supplier, None, f"transaction_header_missing:{field}", source, "error")
        if kind == "snapshot" and not transit_cols:
            self.issue(supplier, None, "transit_headers_missing", source, "error")
        if kind in {"monthly_sales", "history_stock", "snapshot"}:
            self.issue(supplier, None, "warehouse_scope_unconfirmed", source)
        if kind == "history_stock":
            self.issue(supplier, None, "historical_stock_not_confirmed_free_stock", source)
            if not opening:
                self.issue(supplier, None, "historical_stock_snapshot_day_unknown", source)
        if kind == "monthly_sales":
            self.issue(supplier, None, "blank_monthly_cells_preserved_as_nan", source, "info")
        for row in rows:
            def val(i):
                return row[i].value if i is not None and i < len(row) else None
            sku = _text(val(code))
            name = _text(val(name_col))
            if not sku and name and not _label(name).startswith("итого"):
                self.issue(supplier, None, "sku_missing_for_named_row", source, "error")
            if not sku or _label(sku) in SKU_HEADERS or _label(sku).startswith("итого") or _label(name).startswith("итого"):
                continue
            if isinstance(val(code), (int, float)):
                fmt = row[code].number_format
                if re.fullmatch(r"0+", fmt or "") and float(val(code)).is_integer():
                    sku = str(int(val(code))).zfill(len(fmt))
                else:
                    self.issue(supplier, sku, "numeric_sku_leading_zeros_unrecoverable", source)
            values = dict(name=name, unit=_text(val(unit_col)), supplier_article=_text(val(article_col)))
            if kind == "moq":
                col = idx("кратность") if supplier == "Systeme Electric" else idx("мин. разр. к отгр.")
                field = "pack_multiple" if supplier == "Systeme Electric" else "min_order_qty"
                n = self.number(val(col), supplier, sku, source)
                if pd.notna(n) and n <= 0:
                    self.issue(supplier, sku, "nonpositive_order_constraint", source, "error")
                    n = np.nan
                values[field] = n
                values["order_constraint_raw"] = _text(val(col))
                if col is None:
                    self.issue(supplier, None, "order_constraint_header_missing", source, "error")
            if kind == "snapshot":
                values.update(category=_text(val(idx("категория 2026"))))
                for label, field in [("кэф. роста", "growth_coefficient_raw"), ("кэф. сез-ти", "seasonality_coefficient_raw")]:
                    if idx(label) is not None:
                        values[field] = self.number(val(idx(label)), supplier, sku, source)
                if "бухт" in _label(name) and "метр" in _label(name):
                    values["requires_unit_conversion"] = True
                    self.issue(supplier, sku, "coil_meter_conversion_unconfirmed", source, "error")
            self.product(supplier, sku, values, kind, source)
            base = dict(supplier=supplier, sku=sku, source=source)
            if kind == "transactions":
                raw = self.number(val(idx("количество")), supplier, sku, source)
                document = _text(val(idx("документ"))) or ""
                label = _label(document)
                # These exports report sales, not inventory movements. Retain signed
                # outgoing quantities; negative corrections must not become demand.
                if label.startswith("расходная накладная"):
                    quantity = raw
                    tx_type = "return" if quantity < 0 else "sale"
                elif "возврат" in label and ("покупател" in label or "клиент" in label):
                    quantity = -abs(raw)  # Only explicitly identified customer returns.
                    tx_type = "return"
                else:
                    quantity, tx_type = np.nan, "unknown"
                    self.issue(supplier, sku, "unknown_document_type", source, "error")
                when = _date(val(idx("дата")))
                if pd.isna(raw):
                    tx_type = "unknown"
                    self.issue(supplier, sku, "transaction_quantity_missing", source, "error")
                if quantity < 0 and label.startswith("расходная накладная"):
                    self.issue(supplier, None, "negative_invoice_is_return_or_reversal_not_gross_sale", source, "info")
                warehouse = _text(val(idx("склад")))
                if pd.isna(when):
                    self.issue(supplier, sku, "invalid_transaction_date", source, "error")
                if warehouse is None:
                    self.issue(supplier, sku, "unknown_warehouse", source)
                self.rows[kind].append(dict(base, date=when, document_id=_text(val(idx("номер"))),
                    customer_id=None, warehouse=warehouse, quantity=quantity,
                    transaction_type=tx_type, raw_quantity=raw, document=document, unit=values["unit"]))
            elif kind == "monthly_sales":
                for col, month in months.items():
                    self.rows[kind].append(dict(base, month=month, quantity=self.number(val(col), supplier, sku, source),
                        is_complete=month < month_latest))
            elif kind == "history_stock":
                for col, month in months.items():
                    self.rows["stock"].append(dict(base, warehouse=None, as_of=month if opening else pd.NaT,
                        report_month=month, free_stock=np.nan, stock_quantity=self.number(val(col), supplier, sku, source),
                        is_current=False, stock_basis="opening" if opening else "monthly_unspecified", unit=values["unit"]))
            elif kind == "snapshot":
                free = idx("свободный остаток")
                if free is not None:
                    self.rows["stock"].append(dict(base, warehouse=ALL, as_of=cutoff,
                        free_stock=self.number(val(free), supplier, sku, source), is_current=pd.notna(cutoff),
                        stock_basis="free_snapshot", unit=None))
                    if pd.isna(cutoff):
                        self.issue(supplier, sku, "snapshot_date_missing", source, "error")
                for col, h in transit_cols.items():
                    raw = self.number(val(col), supplier, sku, source)
                    if pd.isna(raw):
                        continue  # Blank is not a delivery or a zero; coverage warning below.
                    if raw < 0:
                        self.issue(supplier, sku, "negative_transit_quantity", source, "error")
                    match = re.search(r"поступление до\s*(\d{2}\.\d{2}\.\d{4})", h)
                    when = _date(match[1]) if match else pd.NaT
                    basis = "arrival_deadline" if match else "header_date_unconfirmed"
                    if not match:
                        short = re.search(r"(\d{2})\.(\d{2})(?!\d)", h)
                        if short and pd.notna(cutoff):
                            when = _date(f"{short[1]}.{short[2]}.{cutoff.year}")
                        self.issue(supplier, None, "transit_header_date_requires_confirmation", source)
                    if pd.isna(when):
                        self.issue(supplier, sku, "transit_date_missing", source, "error")
                    conversion = values.get("requires_unit_conversion", False)
                    self.rows["transit"].append(dict(base, warehouse=None, expected_date=when,
                        quantity=np.nan if conversion or raw < 0 else raw, raw_quantity=raw, unit=None,
                        date_basis=basis, shipment=h, requires_unit_conversion=conversion, snapshot_as_of=cutoff))

    def finish(self):
        for key, p in self.products.items():
            supplier, sku = key
            units = self.units[key]
            for field in self.conflicts[key]:
                p[field] = np.nan if field in NUMBER_COLUMNS else None
            if len(units) != 1:
                p["unit"] = None
                self.issue(supplier, sku, "conflicting_units" if units else "unit_missing", "cross_source", "error")
            elif next(iter(units)) not in {"шт", "шт.", "м", "м.", "упак", "упак.", "компл", "компл."}:
                self.issue(supplier, sku, "unit_not_recognized", "cross_source", "error")
                p["unit"] = None
            if len(self.product_sources[key]) == 1:
                self.issue(supplier, sku, "sku_only_in_one_source_type", "cross_source")
            if p.get("requires_unit_conversion"):
                p["min_order_qty"] = np.nan
                p["pack_multiple"] = np.nan
            self.rows["products"].append(p)
        result = {}
        for name, required in SCHEMAS.items():
            df = pd.DataFrame(self.rows[name])
            result[name] = df.reindex(columns=required + [c for c in df.columns if c not in required])
        self.deduplicate(result)
        for supplier in sorted(self.suppliers):
            for missing in ["customer_id", "daily_stockouts", "supplier_lead_times", "bom"]:
                self.issue(supplier, None, f"missing_{missing}", "provided_sources")
            self.issue(supplier, None, "monthly_and_transactions_are_alternative_views_do_not_add", "cross_source", "info")
            self.issue(supplier, None, "transit_blank_cells_do_not_confirm_zero_incoming", "provided_sources")
            self.issue(supplier, None, "order_constraint_units_not_explicit_in_moq", "provided_sources")
            stock = result["stock"]
            if stock.empty or not ((stock.supplier == supplier) & stock.is_current.fillna(False)).any():
                self.issue(supplier, None, "current_stock_missing", "provided_sources", "error")
            for name in ["monthly_sales", "transactions", "transit", "seasonality"]:
                if result[name].empty or not (result[name].supplier == supplier).any():
                    self.issue(supplier, None, f"missing_table_{name}", "provided_sources")
        # Resolve units only through exact supplier/SKU links; never convert quantities.
        for name in ["stock", "transit", "monthly_sales"]:
            df = result[name]
            if not df.empty:
                df["unit"] = [self.products[(s, k)].get("unit") for s,k in zip(df.supplier,df.sku)]
                df["unit_basis"] = "sku_reference"
        for name in ["transactions", "monthly_sales", "stock", "transit"]:
            df = result[name]
            if df.empty:
                continue
            conflict = pd.Series([(s,k) in self.conflicts and "unit" in self.conflicts[(s,k)]
                                  for s,k in zip(df.supplier,df.sku)], index=df.index)
            quantity_col = "free_stock" if name == "stock" else "quantity"
            if conflict.any():
                if "raw_quantity" not in df:
                    df["raw_quantity"] = df[quantity_col]
                df.loc[conflict, quantity_col] = np.nan
        stock = result["stock"]
        current = stock.loc[stock.is_current.fillna(False) & stock.free_stock.notna(), ["supplier", "sku"]]
        current_keys = set(current.itertuples(index=False, name=None))
        monthly_keys = set(zip(result["monthly_sales"].supplier, result["monthly_sales"].sku))
        for key in self.products:
            if key not in current_keys:
                self.issue(*key, "sku_current_free_stock_missing", "cross_source")
            if "snapshot" in self.product_sources[key] and key not in monthly_keys:
                self.issue(*key, "snapshot_sku_missing_monthly_history", "cross_source")
        monthly = result["monthly_sales"]
        for supplier, cutoff in self.cutoffs.items():
            mask = monthly.supplier == supplier
            if mask.any():
                monthly.loc[mask, "is_complete"] = (pd.to_datetime(monthly.loc[mask,"month"]) + pd.offsets.MonthEnd(0)) <= cutoff
        self.reconcile(result)
        quality = pd.DataFrame(self.rows["quality_report"])
        required = SCHEMAS["quality_report"]
        result["quality_report"] = quality.reindex(columns=required + [c for c in quality if c not in required])
        for df in result.values():
            for col in df.columns:
                if col in DATE_COLUMNS:
                    df[col] = pd.to_datetime(df[col], errors="coerce").astype("datetime64[ns]")
                elif col in NUMBER_COLUMNS:
                    df[col] = pd.to_numeric(df[col], errors="coerce").astype(float)
                elif col in {"is_complete", "is_current", "requires_unit_conversion"}:
                    df[col] = df[col].astype("boolean")
                elif col == "month_number":
                    df[col] = df[col].astype("Int64")
                else:
                    df[col] = df[col].astype("string")
        return result

    def deduplicate(self, result):
        """Keep repeated identical facts once; quarantine conflicting keyed facts.

        Transaction line IDs are unavailable: equal lines within one source are
        retained. Repeated documents across sources cannot be safely concatenated.
        """
        stock = result["stock"]
        for supplier, cutoff in self.cutoffs.items():
            if not stock.empty:
                old = (stock.supplier == supplier) & stock.is_current.fillna(False) & (stock.as_of < cutoff)
                if old.any():
                    stock.loc[old, "is_current"] = False
                    self.issue(supplier, None, "older_stock_snapshots_marked_historical", "cross_source", "info")
            transit = result["transit"]
            if not transit.empty:
                old = (transit.supplier == supplier) & transit.snapshot_as_of.notna() & (transit.snapshot_as_of < cutoff)
                if old.any():
                    result["transit"] = transit.loc[~old].reset_index(drop=True)
                    self.issue(supplier, None, "older_transit_snapshots_excluded", "cross_source", "info")
        specs = {
            "monthly_sales": (["supplier", "sku", "month"], "quantity"),
            "stock": (["supplier", "sku", "warehouse", "as_of", "report_month", "stock_basis"], "free_stock"),
            "transit": (["supplier", "sku", "warehouse", "expected_date", "shipment"], "quantity"),
            "seasonality": (["supplier", "category", "month_number"], "factor"),
        }
        for name, (keys, measure) in specs.items():
            df = result[name]
            if df.empty:
                continue
            keys = [k for k in keys if k in df]
            duplicate = df.duplicated(keys, keep=False)
            if not duplicate.any():
                continue
            compare = [c for c in df if c != "source"]
            result[name] = df.drop_duplicates(compare).copy()
            df = result[name]
            conflicting = df.duplicated(keys, keep=False)
            for row in df.loc[df.duplicated(keys, keep=False)].itertuples(index=False):
                self.issue(row.supplier, getattr(row, "sku", None), f"conflicting_{name}_key_excluded", "cross_source", "error")
            # No defensible winner: omit ambiguous rows, preserving a specific issue.
            result[name] = df.loc[~conflicting].reset_index(drop=True)
            for supplier in self.suppliers:
                if ((result[name].supplier == supplier).sum() != (self.rows_count(name, supplier))):
                    self.issue(supplier, None, f"duplicate_{name}_rows_excluded", "cross_source")
        tx = result["transactions"]
        if not tx.empty:
            keys = ["supplier", "sku", "date", "document_id", "warehouse"]
            counts = tx.groupby(keys, dropna=False).source.transform("nunique")
            cross_source = counts > 1
            for supplier, sku in tx.loc[cross_source, ["supplier", "sku"]].drop_duplicates().itertuples(index=False):
                self.issue(supplier, sku, "overlapping_transaction_documents_excluded", "cross_source", "error")
            result["transactions"] = tx.loc[~cross_source].reset_index(drop=True)

    def rows_count(self, name, supplier):
        return sum(r["supplier"] == supplier for r in self.rows[name])

    def reconcile(self, result):
        monthly, transactions = result["monthly_sales"], result["transactions"]
        if monthly.empty:
            return
        monthly["reconciled_transaction_quantity"] = np.nan
        monthly["reconciliation_difference"] = np.nan
        monthly["reconciliation_status"] = "transactions_unavailable"
        if transactions.empty:
            return
        keys = ["supplier", "sku", "month"]
        tx = transactions.assign(month=pd.to_datetime(transactions.date).dt.to_period("M").dt.to_timestamp())
        sums = tx.groupby(keys, dropna=False).quantity.agg(lambda s: s.sum(min_count=len(s))).rename("transaction_quantity").reset_index()
        check = monthly.merge(sums, on=keys, how="outer", indicator=True)
        for (supplier, sku), group in check.loc[check._merge != "both"].groupby(["supplier", "sku"]):
            self.issue(supplier, sku, "sales_history_coverage_differs_between_sources", "cross_source")
        known = check.quantity.notna() & check.transaction_quantity.notna()
        for supplier, sku in check.loc[(check._merge == "both") & ~known, ["supplier", "sku"]].drop_duplicates().itertuples(index=False):
            self.issue(supplier, sku, "reconciliation_incomplete_quantity", "cross_source")
        different = known & ~np.isclose(check.quantity, check.transaction_quantity, rtol=0, atol=1e-6)
        # Numerical agreement is evidence, not approval of warehouse scope or completeness.
        check["reconciliation_status"] = "no_transaction_rows"
        both = check._merge == "both"
        check.loc[both & ~known, "reconciliation_status"] = "incomplete_quantity"
        check.loc[both & known & ~different, "reconciliation_status"] = "matched"
        check.loc[both & different, "reconciliation_status"] = "mismatch"
        check["reconciliation_difference"] = check.quantity - check.transaction_quantity
        evidence = check.loc[check._merge != "right_only", keys + [
            "transaction_quantity", "reconciliation_difference", "reconciliation_status"]].rename(
                columns={"transaction_quantity": "reconciled_transaction_quantity"})
        result["monthly_sales"] = monthly.drop(columns=[
            "reconciled_transaction_quantity", "reconciliation_difference", "reconciliation_status"
        ]).merge(evidence, on=keys, how="left", validate="one_to_one")
        for (supplier, sku), group in check.loc[different].groupby(["supplier", "sku"]):
            months = ",".join(group.month.sort_values().dt.strftime("%Y-%m"))
            self.issue(supplier, sku, f"monthly_transaction_mismatch:{months}", "cross_source", "error",
                       affected_tables="monthly_sales,transactions", affected_months=months,
                       period_start=group.month.min(), period_end=group.month.max() + pd.offsets.MonthEnd(0))
        for supplier in check.supplier.unique():
            subset = check.supplier == supplier
            count = int((known & ~different & subset).sum())
            mismatches = int((different & subset).sum())
            self.issue(supplier, None, f"reconciliation_compared:{count+mismatches};matched:{count};mismatched:{mismatches}", "cross_source", "info")


def load_data(paths: list[str]) -> dict[str, pd.DataFrame]:
    """Read ZIP and/or XLSX paths; return all eight tables, including diagnostics.

    Supplier identification uses the original file or enclosing directory name.
    Invalid inputs are reported in quality_report; programming errors propagate.
    Source workbooks are never modified or extracted to disk.
    """
    loader = _Loader()
    for path in paths:
        p = Path(path)
        try:
            if p.suffix.lower() == ".zip":
                with ZipFile(p) as archive:
                    for member in sorted(archive.namelist()):
                        if member.lower().endswith(".xlsx") and not Path(member).name.startswith("~$"):
                            loader.ingest(f"{p.name}/{member}", archive.read(member))
            elif p.suffix.lower() == ".xlsx":
                loader.ingest(str(p), p.read_bytes())
            else:
                loader.issue(None, None, "unsupported_file_extension", str(p), "error")
        except (OSError, BadZipFile, RuntimeError) as exc:
            loader.issue(_supplier(str(p)), None, f"input_read_error:{type(exc).__name__}", str(p), "error")
    return loader.finish()
