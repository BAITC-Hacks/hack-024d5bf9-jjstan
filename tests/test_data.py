"""Synthetic fixtures only: no commercial workbooks in Git."""
from zipfile import ZipFile
from pathlib import Path
import os

import pandas as pd
import pytest
from openpyxl import Workbook

from data_loader import ALL, SCHEMAS, load_data


def book(tmp_path, name, rows):
    path = tmp_path / name
    wb = Workbook()
    for row in rows:
        wb.active.append(row)
    wb.save(path)
    wb.close()
    return str(path)


def issues(data):
    return set(data['quality_report'].issue)


def test_empty_schema_and_types():
    data = load_data([])
    assert set(data) == set(SCHEMAS)
    for name, columns in SCHEMAS.items():
        assert list(data[name].columns) == columns
    assert str(data['stock'].as_of.dtype) == 'datetime64[ns]'
    assert str(data['transit'].quantity.dtype) == 'float64'


def test_systeme_end_to_end(tmp_path):
    sales = book(tmp_path, 'Ежемесячные продажи Systeme.xlsx', [
        ['Номенклатура', 'Номенклатура.Код', 'янв. 2026', 'сент. 2026', 'Итого'],
        [None, None, 'Количество', 'Количество', 'Количество'],
        ['Тест', '001_', 8, None, 8], ['Итого', None, 8, None, 8]])
    tx = book(tmp_path, 'Динамика Systeme.xlsx', [
        ['Дата', 'Номер', 'Документ', 'Код', 'Номенклатура', 'Ед.', 'Склад', 'Количество'],
        ['10.01.2026 12:30:00', '123', 'Расходная накладная 123', '001_', 'Тест', 'шт', 'Алматы', 10],
        ['12.01.2026', '124', 'Возврат от покупателя 124', '001_', 'Тест', 'шт', 'Алматы', 2]])
    stock = book(tmp_path, 'Товар в пути Systeme на 22.09.2026.xlsx', [
        ['Код 1с', 'Наименование', 'Категория 2026', 'Свободный остаток', 'Зарезервировано', 'Остаток', 'Витрина', 'СЭ в пути 24.09'],
        ['001_', 'Тест', 2, 30, 7, 37, 5, 12]])
    moq = book(tmp_path, 'MOQ Systeme.xlsx', [
        ['Номенклатура.Код', 'Номенклатура', 'Кратность'], ['001_', 'Тест', 6]])
    data = load_data([sales, tx, stock, moq])
    assert data['products'].sku.tolist() == ['001_']
    assert data['products'].pack_multiple.iloc[0] == 6
    assert pd.isna(data['products'].min_order_qty.iloc[0])
    assert data['monthly_sales'].quantity.iloc[0] == 8
    assert pd.isna(data['monthly_sales'].quantity.iloc[1])
    assert data['monthly_sales'].is_complete.tolist() == [True, False]
    assert data['transactions'].quantity.tolist() == [10, -2]
    assert data['transactions'].transaction_type.tolist() == ['sale', 'return']
    assert data['transactions'].customer_id.isna().all()
    assert data['stock'].free_stock.tolist() == [30]
    assert data['stock'].warehouse.tolist() == [ALL]
    assert data['stock'].is_current.all()
    assert data['stock'].as_of.iloc[0] == pd.Timestamp('2026-09-22')
    assert data['transit'].expected_date.iloc[0] == pd.Timestamp('2026-09-24')
    assert not any(i.startswith('monthly_transaction_mismatch') for i in issues(data))


def test_iek_opening_stock_transit_and_moq(tmp_path):
    stock = book(tmp_path, 'Ежемесячные остатки ИЭК.xlsx', [
        ['Номенклатура', 'Ед.', 'Номенклатура.Код', 'сент. 2026', 'Итого'],
        [None, None, None, 'Количество', 'Количество'],
        [None, None, None, 'нач. остаток', 'нач. остаток'],
        ['Тест', 'м', '0002', 100, 100]])
    transit = book(tmp_path, 'Путь ИЭК 22.09.2026.xlsx', [
        ['Код 1с', 'Артикул ИЭК', 'Наименование', 'от 01.09.2026 (поступление до 10.10.2026)'],
        ['0002', 'TEST', 'Кабель ЗАКУПАЮТСЯ БУХТАМИ, САДЯТСЯ МЕТРАЖОМ', 2]])
    moq = book(tmp_path, 'MOQ ИЭК.xlsx', [
        ['Код 1с', 'Наименование', 'Мин. разр. к отгр.'], ['0002', 'Тест', 1], ['0003', 'Тест 2', 5]])
    data = load_data([stock, transit, moq])
    assert len(data['stock']) == 1
    assert not data['stock'].is_current.any()
    assert data['stock'].as_of.iloc[0] == pd.Timestamp('2026-09-01')
    assert data['stock'].stock_quantity.iloc[0] == 100
    assert data['stock'].free_stock.isna().all()
    assert data['transit'].quantity.isna().all()
    assert data['transit'].raw_quantity.iloc[0] == 2
    assert data['transit'].expected_date.iloc[0] == pd.Timestamp('2026-10-10')
    assert data['products'].set_index('sku').loc['0003','min_order_qty'] == 5
    assert data['products'].pack_multiple.isna().all()
    assert 'coil_meter_conversion_unconfirmed' in issues(data)
    assert 'current_stock_missing' in issues(data)


def test_zip_duplicate_and_numeric_format(tmp_path):
    path = book(tmp_path, 'MOQ Systeme.xlsx', [['Код 1с','Кратность'], [12, 4]])
    from openpyxl import load_workbook
    wb=load_workbook(path)
    wb.active['A2'].number_format='00000'
    wb.save(path)
    wb.close()
    archive = tmp_path/'Systeme.zip'
    with ZipFile(archive, 'w') as z:
        z.write(path, 'Systeme/MOQ.xlsx')
    data = load_data([str(archive),path])
    assert data['products'].sku.tolist() == ['00012']
    assert 'duplicate_file_skipped' in issues(data)


def test_unknown_document_never_becomes_sale(tmp_path):
    path = book(tmp_path,'Динамика IEK.xlsx',[
        ['Дата','Номер','Документ','Код','Номенклатура','Ед.','Склад','Количество'],
        ['bad',1,'Перемещение', '01_', 'Тест', None,None,-15],
        ['22.09.2026',2,'Расходная накладная 2','01_', 'Тест','шт',None,-3]])
    data=load_data([path])
    assert pd.isna(data['transactions'].quantity.iloc[0])
    assert data['transactions'].quantity.iloc[1] == -3
    assert data['transactions'].transaction_type.iloc[1] == 'return'
    assert {'unknown_document_type','unknown_warehouse','invalid_transaction_date'} <= issues(data)


def test_missing_file_is_reported(tmp_path):
    data = load_data([str(tmp_path/'missing.zip')])
    assert any(i.startswith('input_read_error') for i in issues(data))
    assert data['products'].empty


def test_monthly_mismatch_reported(tmp_path):
    sales=book(tmp_path,'Ежемесячные продажи IEK.xlsx',[
        ['Номенклатура.Код','янв. 2026'], ['01',99]])
    tx=book(tmp_path,'Динамика IEK.xlsx',[
        ['Дата','Номер','Документ','Код','Количество'],
        ['01.01.2026',1,'Расходная накладная 1','01',3]])
    data=load_data([sales,tx])
    assert 'monthly_transaction_mismatch:2026-01' in issues(data)


def test_seasonality_stops_before_second_block(tmp_path):
    months = ['янв','фев','мар','апр','май','июн','июл','авг','сен','окт','ноя','дек']
    path = book(tmp_path, 'Сезонность ИЭК.xlsx',
                [['Месяц','СЕЗОННОСТЬ']] + [[m,1.0] for m in months] +
                [['Итого',1],[],['Месяц','Другой расчет']] + [[m,None] for m in months])
    data=load_data([path])
    assert len(data['seasonality']) == 12
    assert data['seasonality'].factor.eq(1).all()
    assert data['seasonality'].category.eq(ALL).all()


def test_duplicate_months_without_transactions(tmp_path):
    a=book(tmp_path,'Ежемесячные продажи IEK A.xlsx',[
        ['Код 1с','янв. 2026'],['01',10],['02',3]])
    b=book(tmp_path,'Ежемесячные продажи IEK B.xlsx',[
        ['Код 1с','янв. 2026'],['01',10],['02',4]])
    data=load_data([a,b])
    assert data['monthly_sales'].sku.tolist()==['01']
    assert data['monthly_sales'].quantity.tolist()==[10]
    assert 'conflicting_monthly_sales_key_excluded' in issues(data)


def test_duplicate_stock_and_deliveries_not_added(tmp_path):
    headers=['Код 1с','Наименование','Свободный остаток','СЭ в пути 24.09']
    a=book(tmp_path,'Товар в пути Systeme A 22.09.2026.xlsx',[headers,['01','Тест',3,4]])
    b=book(tmp_path,'Товар в пути Systeme B 22.09.2026.xlsx',[headers,['01','Тест',3,4],[]])
    data=load_data([a,b])
    assert len(data['stock'])==len(data['transit'])==1
    assert data['stock'].free_stock.sum()==3
    assert data['transit'].quantity.sum()==4


def test_unit_conflict_prevents_incompatible_quantities(tmp_path):
    sales=book(tmp_path,'Ежемесячные продажи IEK.xlsx',[
        ['Код 1с','Ед.','янв. 2026'], ['01','шт',12]])
    stock=book(tmp_path,'Ежемесячные остатки IEK.xlsx',[
        ['Код 1с','Ед.','янв. 2026'], ['01','м',20]])
    data=load_data([sales,stock])
    assert data['products'].unit.isna().all()
    assert data['monthly_sales'].quantity.isna().all()
    assert data['monthly_sales'].raw_quantity.tolist()==[12]
    assert 'conflicting_units' in issues(data)


def test_cross_file_transactions_quarantined_but_same_file_lines_kept(tmp_path):
    headers=['Код','Дата','Номер','Документ','Количество','Ед.','Склад']
    row=['01','10.01.2026','A1','Расходная накладная A1',3,'шт','Алматы']
    a=book(tmp_path,'Динамика IEK A.xlsx',[headers,row,row])
    data=load_data([a])
    assert len(data['transactions'])==2
    b=book(tmp_path,'Динамика IEK B.xlsx',[headers,row])
    data=load_data([a,b])
    assert data['transactions'].empty
    assert 'overlapping_transaction_documents_excluded' in issues(data)


def test_snapshot_without_date_not_current(tmp_path):
    path=book(tmp_path,'Товар в пути Systeme.xlsx',[
        ['Код 1с','Свободный остаток','СЭ в пути 24.09'],['01',4,2]])
    data=load_data([path])
    assert not data['stock'].is_current.any()
    assert data['stock'].as_of.isna().all()
    assert data['transit'].expected_date.isna().all()
    assert {'snapshot_date_missing','transit_date_missing'} <= issues(data)


def test_unknown_unit_and_invalid_constraint_reported(tmp_path):
    path=book(tmp_path,'MOQ IEK.xlsx',[
        ['Код 1с','Ед.','Мин. разр. к отгр.'],['01','неизвестно',0],['02','шт','уточнить']])
    data=load_data([path])
    assert data['products'].min_order_qty.isna().all()
    assert {'unit_not_recognized','nonpositive_order_constraint','invalid_numeric_value'} <= issues(data)


def test_no_data_is_inferred_for_missing_capabilities(tmp_path):
    path=book(tmp_path,'MOQ Systeme.xlsx',[['Код 1с','Кратность'],['001',3]])
    data=load_data([path])
    assert data['stockouts'].empty
    assert {'missing_customer_id','missing_daily_stockouts','missing_supplier_lead_times','missing_bom'} <= issues(data)


def test_transaction_nan_is_not_summed_as_zero_in_reconciliation(tmp_path):
    sales=book(tmp_path,'Ежемесячные продажи IEK.xlsx',[
        ['Код 1с','янв. 2026'],['01',3]])
    tx=book(tmp_path,'Динамика IEK.xlsx',[
        ['Код','Дата','Номер','Документ','Количество'],
        ['01','01.01.2026',1,'Расходная накладная 1',3],
        ['01','02.01.2026',2,'Неизвестный документ',10]])
    data=load_data([sales,tx])
    assert 'reconciliation_compared:0;matched:0;mismatched:0' in issues(data)


def test_actual_archives_when_available():
    """Optional local integration check; commercial inputs are never test fixtures."""
    root=os.environ.get('STOCKPILOT_SOURCE_DIR')
    if not root:
        pytest.skip('Set STOCKPILOT_SOURCE_DIR to the private directory containing both ZIPs')
    data=load_data([str(Path(root)/'Systeme electric.zip'),str(Path(root)/'IEK.zip')])
    assert set(data['products'].supplier)=={'Systeme Electric','IEK'}
    assert not data['products'].duplicated(['supplier','sku']).any()
    assert not data['monthly_sales'].duplicated(['supplier','sku','month']).any()
    assert data['seasonality'].groupby('supplier').size().eq(12).all()
    assert data['transit'].expected_date.notna().all()
    assert data['transactions'].customer_id.isna().all()
    assert data['stockouts'].empty
    assert data['products'].sku.str.startswith('0').any()
    assert data['products'].sku.str.endswith('_').any()
    assert data['transactions'].loc[data['transactions'].transaction_type=='sale','quantity'].ge(0).all()
    assert data['transactions'].loc[data['transactions'].transaction_type=='return','quantity'].lt(0).all()
    ieks=data['stock'].loc[data['stock'].supplier=='IEK']
    assert not ieks.is_current.any()
    assert ieks.as_of.dt.day.eq(1).all()
    se=data['stock'].loc[(data['stock'].supplier=='Systeme Electric') & data['stock'].is_current]
    assert not se.empty
    assert se.as_of.eq(pd.Timestamp('2026-09-22')).all()
    september=data['monthly_sales'].loc[data['monthly_sales'].month==pd.Timestamp('2026-09-01')]
    assert not september.is_complete.any()
    assert data['monthly_sales'].quantity.isna().any()


def test_two_snapshots_use_latest_without_double_counting(tmp_path):
    headers=['Код 1с','Свободный остаток','СЭ в пути 24.09']
    older=book(tmp_path,'Товар в пути Systeme 21.09.2026.xlsx',[headers,['01',8,10]])
    newer=book(tmp_path,'Товар в пути Systeme 22.09.2026.xlsx',[headers,['01',12,4]])
    data=load_data([newer,older])
    assert data['stock'].loc[data['stock'].is_current,'free_stock'].tolist()==[12]
    assert data['stock'].loc[~data['stock'].is_current,'free_stock'].tolist()==[8]
    assert data['transit'].quantity.tolist()==[4]


def test_corrupt_xlsx_does_not_discard_good_workbook(tmp_path):
    bad=tmp_path/'MOQ Systeme broken.xlsx'
    bad.write_bytes(b'not a workbook')
    good=book(tmp_path,'MOQ IEK.xlsx',[['Код 1с','Мин. разр. к отгр.'],['01',2]])
    data=load_data([str(bad),good])
    assert data['products'].sku.tolist()==['01']
    assert any(i.startswith('unreadable_workbook') for i in issues(data))


def test_transit_only_iek_still_returns_typed_empty_stock(tmp_path):
    path=book(tmp_path,'Путь ИЭК 22.09.2026.xlsx',[
        ['Код 1с','поступление до 01.10.2026'],['01',10]])
    data=load_data([path])
    assert data['stock'].empty
    assert data['transit'].quantity.tolist()==[10]
    assert 'current_stock_missing' in issues(data)


def test_broken_later_snapshot_does_not_invalidate_valid_stock(tmp_path):
    good=book(tmp_path,'Товар в пути Systeme 22.09.2026.xlsx',[
        ['Код 1с','Свободный остаток','СЭ в пути 24.09'],['01',4,2]])
    bad=tmp_path/'Товар в пути Systeme 23.09.2026.xlsx'
    bad.write_bytes(b'broken')
    data=load_data([good,str(bad)])
    assert data['stock'].is_current.all()
    assert data['transit'].quantity.tolist()==[2]


def test_reconciliation_evidence_preserves_months_and_returns(tmp_path):
    sales=book(tmp_path,'Ежемесячные продажи Systeme.xlsx',[
        ['Код 1с','янв. 2026','февр. 2026','март 2026','апр. 2026','май 2026'],
        ['001_',8,9,4,None,5]])
    tx=book(tmp_path,'Динамика Systeme.xlsx',[
        ['Код','Дата','Номер','Документ','Количество'],
        ['001_','01.01.2026','A','Расходная накладная A',10],
        ['001_','02.01.2026','B','Возврат от покупателя B',2],
        ['001_','01.02.2026','C','Расходная накладная C',3],
        ['001_','01.03.2026','D','Неизвестная операция',4],
        ['001_','01.04.2026','E','Расходная накладная E',2]])
    data=load_data([sales,tx])
    m=data['monthly_sales']
    assert m.reconciliation_status.tolist()==[
        'matched','mismatch','incomplete_quantity','incomplete_quantity','no_transaction_rows']
    assert m.reconciled_transaction_quantity.iloc[:2].tolist()==[8,3]
    assert m.reconciliation_difference.iloc[:2].tolist()==[0,6]
    assert m.reconciliation_difference.iloc[2:].isna().all()
    assert m.quantity.iloc[:3].tolist()==[8,9,4]
    assert pd.isna(m.quantity.iloc[3])
    error=data['quality_report'].loc[data['quality_report'].issue.str.startswith('monthly_transaction_mismatch')].iloc[0]
    assert error.severity=='error'
    assert error.affected_months=='2026-02'
    assert error.period_start==pd.Timestamp('2026-02-01')
    assert error.period_end==pd.Timestamp('2026-02-28')


def test_reconciliation_does_not_match_another_supplier(tmp_path):
    sales=book(tmp_path,'Ежемесячные продажи Systeme.xlsx',[
        ['Код 1с','янв. 2026'],['001_',8]])
    tx=book(tmp_path,'Динамика IEK.xlsx',[
        ['Код','Дата','Номер','Документ','Количество'],
        ['001_','01.01.2026','A','Расходная накладная A',8]])
    data=load_data([sales,tx])
    assert data['monthly_sales'].reconciliation_status.tolist()==['no_transaction_rows']
    assert data['monthly_sales'].reconciled_transaction_quantity.isna().all()
    without_tx=load_data([sales])
    assert without_tx['monthly_sales'].reconciliation_status.tolist()==['transactions_unavailable']
    assert without_tx['monthly_sales'].reconciled_transaction_quantity.isna().all()


def test_reconciliation_reports_exact_noncontiguous_months(tmp_path):
    sales=book(tmp_path,'Ежемесячные продажи IEK.xlsx',[
        ['Код 1с','март 2026','янв. 2026','февр. 2026'],['01',4,4,3]])
    tx=book(tmp_path,'Динамика IEK.xlsx',[
        ['Код','Дата','Номер','Документ','Количество'],
        ['01','01.01.2026','A','Расходная накладная A',3],
        ['01','01.02.2026','B','Расходная накладная B',3],
        ['01','01.03.2026','C','Расходная накладная C',3]])
    data=load_data([sales,tx])
    row=data['quality_report'].loc[data['quality_report'].issue.str.startswith('monthly_transaction_mismatch')].iloc[0]
    assert row.affected_months=='2026-01,2026-03'
    assert row.period_start==pd.Timestamp('2026-01-01')
    assert row.period_end==pd.Timestamp('2026-03-31')
    assert row.issue=='monthly_transaction_mismatch:2026-01,2026-03'
    assert data['monthly_sales'].set_index('month').loc[pd.Timestamp('2026-02-01'),'reconciliation_status']=='matched'


def test_reconciliation_distinguishes_explicit_zero_from_absent_rows(tmp_path):
    sales=book(tmp_path,'Ежемесячные продажи IEK.xlsx',[
        ['Код 1с','янв. 2026'],['01',0],['02',0]])
    tx=book(tmp_path,'Динамика IEK.xlsx',[
        ['Код','Дата','Номер','Документ','Количество'],
        ['01','01.01.2026','A','Расходная накладная A',2],
        ['01','02.01.2026','B','Возврат от покупателя B',2]])
    monthly=load_data([sales,tx])['monthly_sales'].set_index('sku')
    assert monthly.loc['01','reconciliation_status']=='matched'
    assert monthly.loc['01','reconciled_transaction_quantity']==0
    assert monthly.loc['02','reconciliation_status']=='no_transaction_rows'
    assert pd.isna(monthly.loc['02','reconciled_transaction_quantity'])
